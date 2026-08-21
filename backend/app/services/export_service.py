import io
import csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.models.link_item import LinkItem

class ExportService:
    """
    Multi-format Export Service for LinkSense AI.
    Generates Excel (.xlsx), CSV (.csv), and PDF documents with active hyperlinks.
    """

    @classmethod
    def export_to_excel(cls, items: list[LinkItem]) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "Koleksi LinkSense AI"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        data_font = Font(name="Segoe UI", size=10)
        link_font = Font(name="Segoe UI", size=10, color="2563EB", underline="single")
        regular_alignment = Alignment(vertical="center")
        center_alignment = Alignment(horizontal="center", vertical="center")
        wrap_alignment = Alignment(vertical="center", wrap_text=True)

        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        # Headers
        headers = [
            "No",
            "Judul Resource",
            "Platform",
            "Kategori Utama",
            "Sub-Kategori / Use Case",
            "Bahasa",
            "Ringkasan AI",
            "Tags",
            "URL / Tautan",
            "Tanggal Ditambahkan",
        ]

        ws.row_dimensions[1].height = 28
        for col_num, header_title in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header_title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Populate Data Rows
        for row_idx, item in enumerate(items, 2):
            ws.row_dimensions[row_idx].height = 36
            tags_str = ", ".join([t.tag_name for t in item.tags]) if item.tags else ""
            date_str = item.created_at.strftime("%Y-%m-%d %H:%M") if item.created_at else ""
            is_zebra = (row_idx % 2 == 1)

            row_values = [
                (row_idx - 1, center_alignment, data_font, None),
                (item.title, regular_alignment, data_font, None),
                (item.platform, center_alignment, data_font, None),
                (item.primary_category, regular_alignment, data_font, None),
                (item.subcategory or "-", regular_alignment, data_font, None),
                (item.primary_language or "-", center_alignment, data_font, None),
                (item.summary or "-", wrap_alignment, data_font, None),
                (tags_str, regular_alignment, data_font, None),
                (item.url, regular_alignment, link_font, item.url),  # Hyperlink
                (date_str, center_alignment, data_font, None),
            ]

            for col_idx, (val, align, font_style, link_url) in enumerate(row_values, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value = val
                cell.alignment = align
                cell.font = font_style
                cell.border = thin_border
                if is_zebra:
                    cell.fill = zebra_fill
                if link_url:
                    cell.hyperlink = link_url

        # Adjust Column Widths
        column_widths = {
            1: 6,   # No
            2: 32,  # Judul
            3: 15,  # Platform
            4: 24,  # Kategori
            5: 24,  # Subkategori
            6: 12,  # Bahasa
            7: 42,  # Ringkasan
            8: 24,  # Tags
            9: 35,  # URL
            10: 18, # Tanggal
        }
        for col_num, width in column_widths.items():
            ws.column_dimensions[get_column_letter(col_num)].width = width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def export_to_csv(cls, items: list[LinkItem]) -> io.BytesIO:
        output = io.StringIO()
        # UTF-8 with BOM for direct Excel compatibility
        output.write("\ufeff")
        
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
        writer.writerow([
            "No",
            "Judul Resource",
            "Platform",
            "Kategori Utama",
            "Sub-Kategori",
            "Bahasa",
            "Ringkasan AI",
            "Tags",
            "URL",
            "Tanggal Ditambahkan"
        ])

        for idx, item in enumerate(items, 1):
            tags_str = ", ".join([t.tag_name for t in item.tags]) if item.tags else ""
            date_str = item.created_at.strftime("%Y-%m-%d %H:%M") if item.created_at else ""
            writer.writerow([
                idx,
                item.title,
                item.platform,
                item.primary_category,
                item.subcategory or "",
                item.primary_language or "",
                item.summary or "",
                tags_str,
                item.url,
                date_str
            ])

        byte_output = io.BytesIO()
        byte_output.write(output.getvalue().encode("utf-8-sig"))
        byte_output.seek(0)
        return byte_output

    @classmethod
    def export_to_pdf(cls, items: list[LinkItem]) -> io.BytesIO:
        byte_output = io.BytesIO()
        doc = SimpleDocTemplate(
            byte_output,
            pagesize=landscape(letter),
            rightMargin=24,
            leftMargin=24,
            topMargin=24,
            bottomMargin=24,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#0F172A"),
            spaceAfter=4,
        )
        subtitle_style = ParagraphStyle(
            "DocSubtitle",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=14,
        )
        table_header_style = ParagraphStyle(
            "TableHeader",
            parent=styles["Normal"],
            fontSize=9,
            leading=11,
            textColor=colors.white,
            fontName="Helvetica-Bold",
        )
        table_cell_style = ParagraphStyle(
            "TableCell",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#1E293B"),
        )
        table_cell_bold = ParagraphStyle(
            "TableCellBold",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#0F172A"),
            fontName="Helvetica-Bold",
        )
        link_style = ParagraphStyle(
            "TableLink",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#2563EB"),
        )

        elements = []
        # Document Header
        elements.append(Paragraph("Katalog LinkSense AI (Digital Resource Hub)", title_style))
        elements.append(Paragraph(f"Dibuat pada: {datetime.now().strftime('%d %B %Y, %H:%M')} | Total: {len(items)} Tautan", subtitle_style))
        elements.append(Spacer(1, 8))

        # Table Header
        data = [[
            Paragraph("No", table_header_style),
            Paragraph("Judul & Platform", table_header_style),
            Paragraph("Kategori & Sub-Kategori", table_header_style),
            Paragraph("Ringkasan AI", table_header_style),
            Paragraph("Tags", table_header_style),
            Paragraph("Tautan", table_header_style),
        ]]

        for idx, item in enumerate(items, 1):
            tags_str = ", ".join([t.tag_name for t in item.tags]) if item.tags else "-"
            title_p = Paragraph(f"<b>{item.title}</b><br/><font color='#64748B'>[{item.platform}]</font>", table_cell_style)
            cat_p = Paragraph(f"<b>{item.primary_category}</b><br/>{item.subcategory or '-'}", table_cell_style)
            summary_p = Paragraph(item.summary or "-", table_cell_style)
            tags_p = Paragraph(tags_str, table_cell_style)
            link_p = Paragraph(f"<a href='{item.url}' color='#2563EB'><u>Buka Link</u></a>", link_style)

            data.append([
                Paragraph(str(idx), table_cell_style),
                title_p,
                cat_p,
                summary_p,
                tags_p,
                link_p,
            ])

        col_widths = [24, 160, 130, 240, 110, 70]
        pdf_table = Table(data, colWidths=col_widths, repeatRows=1)
        pdf_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ]))

        elements.append(pdf_table)
        doc.build(elements)
        byte_output.seek(0)
        return byte_output

export_service = ExportService()
